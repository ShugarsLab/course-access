#!/usr/bin/env python3
'''
alt_audit.py — triage figure alt text in .pptx decks (especially those exported from Keynote).

Why this exists: Keynote auto-fills PowerPoint's alt-text field with the source
filename ('Screen Shot 2021-03-24 at 8.39.33 PM.png') or a generic label
('Image'). Every presence-based checker — PowerPoint's own and most
automated tools — reports those slides as fully compliant. This script checks
whether alt text is *useful*, not whether it exists, and separates figures that
carry content from decorative/repeated images that do not.

Usage:
    python3 alt_audit.py deck.pptx [deck2.pptx ...]
    python3 alt_audit.py --csv findings.csv deck.pptx
'''

import argparse, csv, hashlib, os, re, sys, collections

try:
    from pptx import Presentation
except ImportError:
    sys.exit('python-pptx required:  pip3 install python-pptx')

FILE_EXT   = re.compile(r'\.(png|jpe?g|gif|tiff?|pdf|mov|svg|heic|emf|wmf)$', re.I)
SCREENSHOT = re.compile(r'^(screen\s?shot|screenshot|pasted)', re.I)
GENERIC    = re.compile(
    r'^(image|picture|graphic|photo|shape|line|object|content placeholder'
    r'|pasted[- ]?(image|movie|file))([ _-]?\d+)?$', re.I)
REPEATWORD = re.compile(r'^(\w+)( \1)+$', re.I)      # e.g. 'Line Line'

# A figure smaller than this fraction of the slide is assumed decorative.
SIZE_FLOOR = 0.03
# An identical image appearing more than this many times is assumed branding.
REPEAT_CEIL = 3
# Shorter than this is too terse to describe a quantitative figure.
MIN_USEFUL_LEN = 15


def classify(alt, shape_name):
    '''Return one of MISSING / FILENAME / GENERIC / ECHOES_NAME / TOO_SHORT / SUBSTANTIVE.'''
    alt = (alt or '').strip()
    if not alt:
        return 'MISSING'
    if FILE_EXT.search(alt) or SCREENSHOT.match(alt):
        return 'FILENAME'
    if GENERIC.match(alt) or REPEATWORD.match(alt):
        return 'GENERIC'
    if alt.lower() == (shape_name or '').strip().lower():
        return 'ECHOES_NAME'
    if len(alt) < MIN_USEFUL_LEN:
        return 'TOO_SHORT'
    return 'SUBSTANTIVE'


def alt_of(shape):
    try:
        return shape._element._nvXxPr.cNvPr.get('descr')
    except Exception:
        return None


def walk(shapes):
    '''Yield every shape, descending into groups (builds are usually groups).'''
    for shape in shapes:
        yield shape
        if str(shape.shape_type).startswith('GROUP'):
            try:
                yield from walk(shape.shapes)
            except Exception:
                pass


def audit(path):
    prs = Presentation(path)
    slide_area = prs.slide_width * prs.slide_height
    records, hashes = [], collections.Counter()

    for index, slide in enumerate(prs.slides, 1):
        for shape in walk(slide.shapes):
            if 'PICTURE' not in str(shape.shape_type):
                continue
            try:
                digest = hashlib.md5(shape.image.blob).hexdigest()[:8]
            except Exception:
                digest = '?'
            hashes[digest] += 1
            try:
                size = (shape.width * shape.height) / slide_area
            except Exception:
                size = 0.0
            records.append({
                'slide': index,
                'shape': shape.name,
                'hash': digest,
                'size': size,
                'alt': (alt_of(shape) or '').strip(),
                'verdict': classify(alt_of(shape), shape.name),
            })

    for record in records:
        decorative = record['size'] < SIZE_FLOOR or hashes[record['hash']] > REPEAT_CEIL
        record['decorative'] = decorative
        if record['verdict'] == 'SUBSTANTIVE':
            record['action'] = 'ok'
        elif decorative:
            record['action'] = 'mark-decorative'
        else:
            record['action'] = 'NEEDS-DESCRIPTION'
    return records


def report(path, records):
    counts = collections.Counter(r['verdict'] for r in records)
    todo = [r for r in records if r['action'] == 'NEEDS-DESCRIPTION']
    decor = [r for r in records if r['action'] == 'mark-decorative']

    print(f'\n{os.path.basename(path)}')
    print(f'  {len(records)} images: '
          f"{counts.get('SUBSTANTIVE', 0)} described, "
          f"{counts.get('FILENAME', 0)} filename, "
          f"{counts.get('GENERIC', 0)} generic, "
          f"{counts.get('MISSING', 0)} empty")
    print(f'  -> {len(todo)} need a real description; '
          f'{len(decor)} should be marked decorative')

    if todo:
        print('\n  WORKLIST')
        by_slide = collections.defaultdict(list)
        for record in todo:
            by_slide[record['slide']].append(record)
        for slide in sorted(by_slide):
            for record in by_slide[slide]:
                current = record['alt'] or '(empty)'
                print(f"    slide {slide:>3}  {record['size']*100:>4.0f}% of slide  "
                      f"currently: {current[:58]!r}")


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('decks', nargs='+')
    parser.add_argument('--csv', help='write all findings to this CSV')
    args = parser.parse_args()

    everything = []
    for deck in args.decks:
        records = audit(deck)
        for record in records:
            record['deck'] = os.path.basename(deck)
        report(deck, records)
        everything.extend(records)

    if args.csv:
        with open(args.csv, 'w', newline='') as handle:
            fields = ['deck', 'slide', 'shape', 'size', 'verdict', 'decorative', 'action', 'alt']
            writer = csv.DictWriter(handle, fieldnames=fields, extrasaction='ignore')
            writer.writeheader()
            writer.writerows(everything)
        print(f'\nWrote {len(everything)} rows to {args.csv}')


if __name__ == '__main__':
    main()
