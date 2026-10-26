"""pptx2notes.py: turn a remediated .pptx into a text-only companion document (.docx).

Usage: python3 pptx2notes.py <deck.pptx> [-o out.docx] [--changes changes.csv]

Writes the Word version under the deck's own name, <deck>.docx, beside the deck, and the
as-generated copy to _snapshots/<deck>_generated.docx (a folder created beside the deck). If
<deck>.docx already exists it is treated as the instructor's edited master and is NOT
overwritten: only the generated copy is written (compare the two with Word's Review > Compare).
With -o, writes only that file.

The converter does not write descriptions. It reads the remediation already in the deck --
alt text, decorative flags, reading order, equations, tables -- and turns it into a document
read like lecture notes. Each slide is read in reading order (the order PowerPoint's Reading
Order pane shows), and the converter writes:
- a heading for every slide, numbered as the slide's page in the exported PDF:
  "Slide 37: <the title as on the slide>", repeated titles included; a build that only adds to
  the slide before it shares one heading, "Slides 37 to 40: <title>" (SLIDE_HEADINGS,
  BUILD_RANGES); hidden slides are skipped and not counted, as in PowerPoint's PDF export;
- Heading 2 at section breaks, taken from repeated agenda slides (AGENDA_RE); slide headings
  are Heading 3 below them (Heading 2 in a deck with no section breaks);
- slide text in reading order; bulleted paragraphs as real Word bulleted lists, nested as on the
  slide (REAL_BULLETS); numbered steps (PowerPoint's autonumbering) as a Word numbered list,
  numbered as on the slide;
- quote marks removed, and colons in sentences turned into commas or dropped (COLONS_TO_COMMAS);
  colons inside names (Gr_Liv_Area:Bldg_Type), before "=" (:=), standing alone, and in headings
  are kept; code is never changed;
- code and output in the Code style, verbatim, curly quotes straightened, each block opened
  with "Begin R code block." and, when plain text follows, closed with "End R code block."
  (not before a figure, an equation, a heading or a table) (CODE_LANGUAGE);
- R output that prints as a table (summary() coefficients, aov(), anova(), lavaan estimates, a
  named vector) as a real Word table after the line "R output as a table", with significance
  stars in their own column and Pr(>|t|) and similar headed p-value (R_OUTPUT_TABLES);
- in prose: a superscript number as words ("R-squared", "0.95 to the power 10"), a subscript as
  "sub" ("beta sub k"), and the times sign as "times" (SUPERSCRIPTS_AS_WORDS,
  SUBSCRIPTS_AS_WORDS, TIMES_AS_WORD);
- numbers of 1,000 or more outside code, equations, and R output tables: "140 thousand" when
  round, words when exact ("two thousand four hundred thirteen"), digits without the comma for
  decimals and table cells; years are left alone (NUMBER_RULE). Every change is listed;
- short inline equations (a symbol or two inside a sentence) as words, because Word announces
  "begin equation" for each; full equations stay live Word equations;
- two equations split by a lone operator in the text merged into one equation;
- a group made only of equation pieces (for example argmin over a sum, with a "For ..." piece)
  as ONE live equation, not as its figure description (flagged for review);
- a short symbol followed by its meaning joined as one line ("beta sub 0 means intercept"),
  only when the meaning sits right beside the symbol and is a phrase, not another equation;
- a figure (picture or group) as its alt text, in place, prefixed "Figure.";
- a narrated arrow's alt text joined to the item it leads into; a repeated item that the
  narration continues from is kept;
- tables as Word tables with a repeating header row;
- space after code blocks and tables as paragraph spacing, not empty paragraphs (a screen reader
  announces an empty paragraph as "blank");
- the title slide's other lines (course name, instructor, pronouns) under the document title;
- the document title and language in the file's properties (LANGUAGE);
- equation-shaped figures as live equations: a hand-built fraction (pieces above and below a drawn
  line) and rows of equations whose gaps hold named icons;
- Keynote's line separator (U+2028) as a space; line breaks inside a prose paragraph as spaces
  (SOFT_BREAKS_AS_SPACES); a sentence Keynote split over two paragraphs of one text box joined
  back, unless the second is a bulleted item (JOIN_FRAGMENTS); a slide's own "In R:" label
  before code dropped, since the Begin marker replaces it.
Skips decorative objects, slide numbers, speaker notes, and any paragraph identical to one on the
previous slide (repeats in a build); a code block is skipped only when the whole block repeats,
never line by line. Also: a line starting with a relation sign ("= 42") joins the line before it
in the same box; a short label inside a described figure, its words already in the description,
is left to the description; a label ending in a colon that sits just above a text box but is read
after it leads into that text; a group of bare arrows with alt text is narration, like a single
arrow, and narration with nothing after it on the slide goes before the item nearest the arrow; a
wordless annotation beside a table row goes into that row's last cell. Lone labels from
undescribed groups and narration joined to a lone symbol are flagged. Prints a review list of
anything to check, and every number and symbol it rewrote (before, after, where); --changes
also writes that list as a CSV file.

SETTINGS YOU MAY WANT TO CHANGE
These encode one course's choices (its "course preferences"); another course or another listener
may prefer otherwise. Edit them in this file (there is no settings file yet); search for each name.
- SLIDE_HEADINGS, SLIDE_HEADING_FMT, BUILD_RANGES, BUILD_HEADING_FMT: a numbered heading for
  every slide, and one heading for a build. Off: slide titles as headings only when they change,
  without numbers.
- REAL_BULLETS: off writes bullets as indented paragraphs.
- NUMBER_RULE, NUMBER_WORD_HYPHENS: the large-number wording; hyphens off writes "forty eight".
- R_OUTPUT_TABLES; SUPERSCRIPTS_AS_WORDS, SUBSCRIPTS_AS_WORDS, TIMES_AS_WORD.
- COLONS_TO_COMMAS, and QUOTES and CODE_QUOTES with clean_chars below them: quote marks are
  removed from text; code keeps straight quotes.
- AGENDA_RE: the slide title that marks an agenda slide ("Today's agenda"). Repeated agenda
  slides become the document's section headings.
- Document title fallback (in main): with no Title property set on the deck, the title slide's
  line starting "Week N" becomes the document title.
- CODE_LANGUAGE: the language named in the code markers ("Begin R code block."); IN_R_RE drops a
  slide's own "In R:" label.
- LANGUAGE: the document's language, as a language tag (en-US).
- MONO: the fonts that mark text as code. BODY_FONT, CODE_FONT, CODE_SIZE: the document's fonts.
- FIGURE_PREFIX ("Figure. ") before each description; " means " for a symbol joined to its
  meaning (in Notes.para).
- SHORT_MATH_CHARS: an inline equation this short, with no structure, is written as words.
- SOFT_BREAKS_AS_SPACES, JOIN_FRAGMENTS: joining lines and sentence fragments in prose.
- LEADS_ON, GREEK_NAMES, and GREEK_LETTERS: English word lists used when joining fragments and
  writing symbols as words.
- _snapshots/ (in main): where the as-generated copy is written.
"""
import argparse, copy, difflib, os, re, shutil, sys
from lxml import etree
from pptx import Presentation
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mathspeech

NS = {
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
    'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
    'a14': 'http://schemas.microsoft.com/office/drawing/2010/main',
    'adec': 'http://schemas.microsoft.com/office/drawing/2017/decorative',
}
def q(pfx, tag): return '{%s}%s' % (NS[pfx], tag)
MONO = {'menlo', 'courier', 'courier new', 'consolas', 'monaco', 'source code pro',
        'andale mono', 'fira mono', 'fira code', 'sf mono', 'lucida console', 'dejavu sans mono'}
AGENDA_RE = re.compile(r'^\s*today.?s agenda\s*$', re.I)
SKIP_PH = {'sldNum', 'dt', 'ftr', 'hdr'}
RELATION_RE = re.compile(r'[=≠<>≤≥≈~∼]')
OPERATOR_TEXT = set('  -−+=×·/,')
# structures that make an equation more than "a symbol or two"
MATH_STRUCT = {'f', 'nary', 'rad', 'd', 'func', 'eqArr', 'limLow', 'limUpp', 'groupChr', 'm',
               'box', 'borderBox', 'sPre', 'phant'}
SHORT_MATH_CHARS = 6
SOFT_BREAKS_AS_SPACES = True
JOIN_FRAGMENTS = True         # a sentence Keynote split over two paragraphs becomes one
IN_R_RE = re.compile(r'^\s*in r\s*[,:]?\s*$', re.I)
GROUP_TOKEN_RE = re.compile(r'^(?:argmin|argmax|min|max|for|lim|[\s()\[\],=+\-−·×/0-9])*$', re.I)
EMU_PT = 12700
NARRATION_GEOMS = {'line', 'straightConnector1', 'bentConnector2', 'bentConnector3', 'curvedConnector2',
                   'curvedConnector3', 'rightArrow', 'leftArrow', 'upArrow', 'downArrow',
                   'leftBrace', 'rightBrace', 'leftBracket', 'rightBracket'}
RELATION_START_RE = re.compile(r'^\s*[=<>≤≥≈≠]')
# --- the companion document's format: one course's preferences (see SETTINGS in the docstring)
SLIDE_HEADINGS = True               # a numbered heading on every slide
SLIDE_HEADING_FMT = 'Slide {n}: {title}'
BUILD_RANGES = True                 # a build that only adds to the slide before shares its heading
BUILD_HEADING_FMT = 'Slides {first} to {last}: {title}'
REAL_BULLETS = True                 # off: bullets become indented paragraphs
COLONS_TO_COMMAS = True             # colons in sentences become commas, or are dropped at the end
NUMBER_RULE = True                  # numbers of 1,000 or more as words, "thousand", or digits
NUMBER_WORD_HYPHENS = False         # True: "forty-eight"; False: "forty eight"
R_OUTPUT_TABLES = True              # table-like R output becomes a real Word table
SUPERSCRIPTS_AS_WORDS = True        # R2 -> R-squared, 10 to the power 4, in prose
SUBSCRIPTS_AS_WORDS = True          # beta with subscript k -> beta sub k, in prose
TIMES_AS_WORD = True                # a × b -> a times b, in prose
CODE_LANGUAGE = 'R'                 # "Begin R code block."
LANGUAGE = 'en-US'                  # the document's language tag
BODY_FONT, CODE_FONT, CODE_SIZE = 'Times New Roman', 'Consolas', 10.5
FIGURE_PREFIX = 'Figure. '
BULLET_STYLES = ('List Bullet', 'List Bullet 2', 'List Bullet 3')
GREEK_LETTERS = {'α': 'alpha', 'β': 'beta', 'γ': 'gamma', 'δ': 'delta', 'ε': 'epsilon',
                 'θ': 'theta', 'λ': 'lambda', 'μ': 'mu', 'π': 'pi', 'ρ': 'rho', 'σ': 'sigma',
                 'τ': 'tau', 'φ': 'phi', 'χ': 'chi', 'ψ': 'psi', 'ω': 'omega', 'η': 'eta'}
AUTONUM_FMT = {   # PowerPoint buAutoNum type -> Word numFmt, lvlText
    'arabicPeriod': ('decimal', '%1.'), 'arabicParenR': ('decimal', '%1)'),
    'arabicParenBoth': ('decimal', '(%1)'), 'arabicPlain': ('decimal', '%1'),
    'alphaLcPeriod': ('lowerLetter', '%1.'), 'alphaLcParenR': ('lowerLetter', '%1)'),
    'alphaLcParenBoth': ('lowerLetter', '(%1)'), 'alphaUcPeriod': ('upperLetter', '%1.'),
    'alphaUcParenR': ('upperLetter', '%1)'), 'alphaUcParenBoth': ('upperLetter', '(%1)'),
    'romanLcPeriod': ('lowerRoman', '%1.'), 'romanLcParenR': ('lowerRoman', '%1)'),
    'romanUcPeriod': ('upperRoman', '%1.'), 'romanUcParenR': ('upperRoman', '%1)'),
}
GENERIC_NAME_RE = re.compile(r'^(image|picture|graphic|group|shape|freeform|rectangle|oval|line|object|icon)?\s*\d*$', re.I)

review = []
QUOTES = set('“”"‘')
CODE_QUOTES = {'‘': "'", '’': "'", '“': '"', '”': '"'}

def local(e): return etree.QName(e).localname

def clean_chars(chars, colons=None):
    """chars: list of (ch, seg). Drop quote marks; turn colons into commas or drop them.
    Keeps apostrophes (a right single quote between two letters), colons between digits, colons
    inside names (Gr_Liv_Area:Bldg_Type) or before "=" (:=), and a colon standing alone ("joined
    by :")."""
    if colons is None: colons = COLONS_TO_COMMAS
    out = []
    n = len(chars)
    for i, (c, seg) in enumerate(chars):
        prev = chars[i - 1][0] if i > 0 else ''
        nxt = chars[i + 1][0] if i + 1 < n else ''
        if c in QUOTES:
            continue
        if c == '’' and not (prev.isalpha() and nxt.isalpha()):
            continue
        if c == ':':
            if not colons:
                out.append((c, seg)); continue
            if prev.isdigit() and nxt.isdigit():
                out.append((c, seg)); continue
            if (prev and nxt and not prev.isspace() and not nxt.isspace()) or nxt == '=':
                out.append((c, seg)); continue           # inside a name, or :=
            if not prev or prev.isspace():
                out.append((c, seg)); continue           # a colon standing alone is the symbol itself
            rest = ''.join(ch for ch, _ in chars[i + 1:]).split('\n')[0]
            if not rest.strip():
                stats_inc('colons removed'); continue        # trailing colon
            stats_inc('colons to commas')
            out.append((',', seg))
            if nxt not in (' ', '\n'): out.append((' ', seg))
            continue
        out.append((c, seg))
    return out

def clean_text(t, colons=None):
    t = t.replace(' ', ' ')
    return ''.join(c for c, _ in clean_chars([(c, 0) for c in t], colons)).strip()

def clean_title(t):
    """A slide title for its heading: as on the slide, quote marks removed, colons kept."""
    return ' '.join(clean_text(t, colons=False).split())

def stats_inc(k, n=1):
    if notes is not None: notes.stats[k] = notes.stats.get(k, 0) + n

def lower_first(t):
    if len(t) > 1 and t[0].isupper() and t[1].islower(): return t[0].lower() + t[1:]
    return t

def is_decorative(el):
    cnv = el.find('.//' + q('p', 'cNvPr'))
    if cnv is None: return False
    for d in cnv.iter(q('adec', 'decorative')):
        if d.get('val') == '1': return True
    return False

def descr(el):
    cnv = el.find('.//' + q('p', 'cNvPr'))
    return (cnv.get('descr') or '').strip() if cnv is not None else ''

def ph_info(sp):
    ph = sp.find('.//' + q('p', 'ph'))
    if ph is None: return None
    return ph.get('type', 'body'), ph.get('idx', '0')

def unwrap(child):
    """mc:AlternateContent -> the element inside mc:Choice."""
    if child.tag == q('mc', 'AlternateContent'):
        ch = child.find(q('mc', 'Choice'))
        if ch is not None and len(ch): return ch[0]
        fb = child.find(q('mc', 'Fallback'))
        if fb is not None and len(fb): return fb[0]
        return None
    return child

def clean_math(el):
    """Copy OMML out of PowerPoint, dropping DrawingML run properties Word rejects."""
    m = copy.deepcopy(el)
    for bad in list(m.iter('{%s}*' % NS['a'])):
        par = bad.getparent()
        if par is not None: par.remove(bad)
    return m

def M(tag): return etree.Element(q('m', tag), nsmap={'m': NS['m']})

def mrun(text, plain=False):
    r = M('r')
    if plain:
        rp = etree.SubElement(r, q('m', 'rPr')); s = etree.SubElement(rp, q('m', 'sty')); s.set(q('m', 'val'), 'p')
    t = etree.SubElement(r, q('m', 't')); t.text = text
    return r

# ------------------------------------------------------------------ geometry
def local_box(el):
    for path in (q('p', 'spPr'), q('p', 'grpSpPr'), q('p', 'xfrm')):
        c = el.find(path)
        if c is None: continue
        x = c if c.tag == q('p', 'xfrm') else c.find(q('a', 'xfrm'))
        if x is None: continue
        off, ext = x.find(q('a', 'off')), x.find(q('a', 'ext'))
        if off is None or ext is None: continue
        return [int(off.get('x')), int(off.get('y')), int(ext.get('cx')), int(ext.get('cy'))]
    return None

def abs_box(el):
    b = local_box(el)
    if b is None: return None
    x, y, w, h = b
    anc = el.getparent()
    while anc is not None:
        if anc.tag == q('p', 'grpSp'):
            gx = anc.find(q('p', 'grpSpPr') + '/' + q('a', 'xfrm'))
            if gx is not None:
                off, ext = gx.find(q('a', 'off')), gx.find(q('a', 'ext'))
                co, ce = gx.find(q('a', 'chOff')), gx.find(q('a', 'chExt'))
                if None not in (off, ext, co, ce):
                    sx = int(ext.get('cx')) / int(ce.get('cx')) if int(ce.get('cx')) else 1
                    sy = int(ext.get('cy')) / int(ce.get('cy')) if int(ce.get('cy')) else 1
                    x = int(off.get('x')) + (x - int(co.get('x'))) * sx; w *= sx
                    y = int(off.get('y')) + (y - int(co.get('y'))) * sy; h *= sy
        anc = anc.getparent()
    return [x, y, w, h]

def beside(sym, other):
    """Is `other` laid out as the meaning of the symbol at `sym` (to its right, or just below)?"""
    if sym is None or other is None: return False
    sx, sy, sw, sh = sym; ox, oy, ow, oh = other
    v_overlap = sy < oy + oh and oy < sy + sh
    gap_right = ox - (sx + sw)
    if v_overlap and -0.5 * sw <= gap_right <= 150 * EMU_PT: return True
    gap_below = oy - (sy + sh)
    return -10 * EMU_PT <= gap_below <= 60 * EMU_PT and abs(ox - sx) <= 150 * EMU_PT

# ------------------------------------------------------------------ fonts
def latin_of(container):
    if container is None: return None
    l = container.find(q('a', 'latin'))
    return l.get('typeface') if l is not None else None

def run_face(r, ap, tb):
    f = latin_of(r.find(q('a', 'rPr')))
    if f: return f
    ppr = ap.find(q('a', 'pPr'))
    if ppr is not None:
        f = latin_of(ppr.find(q('a', 'defRPr')))
        if f: return f
    lvl = int(ppr.get('lvl', '0')) if ppr is not None else 0
    lst = tb.find(q('a', 'lstStyle')) if tb is not None else None
    if lst is not None:
        lp = lst.find(q('a', 'lvl%dpPr' % (lvl + 1)))
        if lp is not None:
            f = latin_of(lp.find(q('a', 'defRPr')))
            if f: return f
    return ''

def is_mono(ap, tb=None):
    """Code if every run with visible text resolves (through the paragraph and the text box's
    list style) to a monospace face. A prose bullet with one `lm()` run is not code."""
    faces = [run_face(r, ap, tb).lower() for r in ap.iter(q('a', 'r'))
             if ''.join(x.text or '' for x in r.iter(q('a', 't'))).strip()]
    return bool(faces) and all(f in MONO for f in faces)

# ------------------------------------------------------------------ maths
def math_text(el):
    return ''.join(x.text or '' for x in el.iter(q('m', 't')))

def is_short_math(om):
    lin = math_text(om).replace(' ', '')
    if not lin or len(lin) > SHORT_MATH_CHARS: return False
    return not any(local(e) in MATH_STRUCT for e in om.iter()
                   if isinstance(e.tag, str) and e.tag.startswith('{%s}' % NS['m']))

def math_words(om, where):
    s = mathspeech.speak([om]).strip()
    if 'is distributed as' in s:
        review.append(f'{where}: "{math_text(om)}" written as "{s}"; in an R formula say "tilde"')
    return s

def strip_trailing_colon(math_el):
    ts = [x for x in math_el.iter(q('m', 't')) if (x.text or '').strip()]
    if ts and ts[-1].text.rstrip().endswith(':'):
        ts[-1].text = ts[-1].text.rstrip()[:-1]

# ------------------------------------------------------------------ document
class Notes:
    def __init__(self, title):
        self.doc = Document()
        st = self.doc.styles
        for name in ('Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Bullet 2', 'List Bullet 3', 'Title'):
            s = st[name]; s.font.name = BODY_FONT
            rpr = s.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
            if rf is None: rf = etree.SubElement(rpr, qn('w:rFonts'))
            for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(k), BODY_FONT)
            rf.attrib.pop(qn('w:asciiTheme'), None); rf.attrib.pop(qn('w:hAnsiTheme'), None)
            if name.startswith('Heading') or name == 'Title':
                s.font.color.rgb = None
        st['Normal'].font.size = Pt(12)
        code = st.add_style('Code', WD_STYLE_TYPE.PARAGRAPH); code.base_style = st['Normal']
        code.font.name = CODE_FONT; code.font.size = Pt(CODE_SIZE)
        code.paragraph_format.space_after = Pt(0)
        self.code_style_id = code.style_id
        self.doc.core_properties.title = title
        self.doc.core_properties.language = LANGUAGE
        dd = st.element.find(qn('w:docDefaults'))           # the language every run inherits
        rpr = dd.find(qn('w:rPrDefault') + '/' + qn('w:rPr')) if dd is not None else None
        if rpr is not None:
            lang = rpr.find(qn('w:lang'))
            if lang is None: lang = etree.SubElement(rpr, qn('w:lang'))
            lang.set(qn('w:val'), LANGUAGE)
        self.changes = []        # (where, kind, before, after): every number and symbol rewritten
        self.slide = None        # the python-pptx slide being written, for inherited bullets
        self.r_tables = set()    # Word tables made from R output: the number rule leaves them alone
        self.doc.add_heading(title, level=1)
        self.prefix = ''
        self.pending_symbol = None; self.pending_box = None
        self.prev_slide = set(); self.this_slide = set()
        self.last_title = None; self.section_heading = None
        self.last_skipped = None
        self.prev_tables = []; self.this_slide_tables = []; self.where = ''
        self.stats = {'paragraphs': 0, 'figures': 0, 'equations': 0, 'tables': 0, 'repeats skipped': 0}
        self.num_last = {}; self.num_seq = 0

    def new_num(self, typ, start):
        """A fresh one-level Word list (its own abstractNum), starting at `start`."""
        fmt, text = AUTONUM_FMT.get(typ, ('decimal', '%1.'))
        numbering = self.doc.part.numbering_part.element
        W = lambda t: qn('w:' + t)
        ids = [int(e.get(W('abstractNumId'))) for e in numbering.findall(W('abstractNum'))]
        nids = [int(e.get(W('numId'))) for e in numbering.findall(W('num'))]
        aid, nid = max(ids + [0]) + 1, max(nids + [0]) + 1
        an = etree.Element(W('abstractNum')); an.set(W('abstractNumId'), str(aid))
        etree.SubElement(an, W('multiLevelType')).set(W('val'), 'singleLevel')
        lvl = etree.SubElement(an, W('lvl')); lvl.set(W('ilvl'), '0')
        etree.SubElement(lvl, W('start')).set(W('val'), str(start))
        etree.SubElement(lvl, W('numFmt')).set(W('val'), fmt)
        etree.SubElement(lvl, W('lvlText')).set(W('val'), text)
        etree.SubElement(lvl, W('lvlJc')).set(W('val'), 'left')
        ind = etree.SubElement(etree.SubElement(lvl, W('pPr')), W('ind'))
        ind.set(W('left'), '360'); ind.set(W('hanging'), '360')
        first_num = numbering.find(W('num'))
        if first_num is not None: first_num.addprevious(an)
        else: numbering.append(an)
        num = etree.SubElement(numbering, W('num')); num.set(W('numId'), str(nid))
        etree.SubElement(num, W('abstractNumId')).set(W('val'), str(aid))
        return nid

    def number(self, p, typ, n, seq, lvl):
        """Put Word paragraph p in the numbered list for PowerPoint sequence `seq`, as item n.
        A gap (an item skipped as a repeat in a build) starts a new list at n, so the number
        heard is always the number on the slide."""
        last = self.num_last.get(seq)
        nid = last[0] if last is not None and last[1] == n - 1 else self.new_num(typ, n)
        self.num_last[seq] = (nid, n)
        ppr = p._p.get_or_add_pPr(); numpr = ppr.get_or_add_numPr()
        numpr.get_or_add_ilvl().val = 0; numpr.get_or_add_numId().val = nid
        p.paragraph_format.left_indent = Pt(18 * min(lvl, 3) + 18)
        p.paragraph_format.first_line_indent = Pt(-18)
        stats_inc('numbered list items')

    def para(self, style=None, key=None, force=False, redo=None):
        if key is not None:
            if key in self.prev_slide and not force:
                self.stats['repeats skipped'] += 1
                self.this_slide.add(key)
                self.last_skipped = redo
                return None
            self.this_slide.add(key)
        self.last_skipped = None
        p = self.doc.add_paragraph(style=style)
        if self.pending_symbol is not None:
            ms = self.pending_symbol; self.pending_symbol = None; self.pending_box = None
            if all(is_short_math(m) for m in ms):
                p.add_run(' '.join(math_words(m, self.where) for m in ms) + ' means ')
                stats_inc('inline equations to words', len(ms))
            else:
                for m in ms: p._p.append(m)
                p.add_run(' means ')
            stats_inc('symbol keys joined')
        if self.prefix:
            p.add_run(self.prefix + ' '); self.prefix = ''
            if hasattr(self, 'prefixed'): self.prefixed.add(p._p)
        self.stats['paragraphs'] += 1
        return p

    def new_slide(self, where=''):
        self.prev_slide = self.this_slide; self.this_slide = set()
        self.prev_tables = self.this_slide_tables; self.this_slide_tables = []
        self.where = where; self.last_skipped = None
        self.slide_shapes = []      # (box, first Word paragraph) per text shape written on this slide
        self.slide_tables = []      # (box, row heights, Word table) per table written on this slide
        self.fig_boxes = []         # (box, description words) per described figure on this slide
        self.lead_labels = set()    # labels moved before the text below them
        self.prefix_box = None; self.prefix_points = []
        self.slide_items = []       # (box, first Word paragraph or None for a table), in writing order
        self.prefixed = set()       # Word paragraphs that already took a narration lead-in

notes = None

def para_key(ap):
    t = ''.join(x.text or '' for x in ap.iter(q('a', 't')))
    mt = ''.join(x.text or '' for x in ap.iter(q('m', 't')))
    return re.sub(r'\s+', ' ', (t + '|' + mt)).strip()

def bullet_level(ap, placeholder_body):
    ppr = ap.find(q('a', 'pPr'))
    lvl = int(ppr.get('lvl', '0')) if ppr is not None else 0
    if ppr is not None and ppr.find(q('a', 'buNone')) is not None: return None
    has_bu = ppr is not None and (ppr.find(q('a', 'buChar')) is not None or ppr.find(q('a', 'buAutoNum')) is not None)
    if placeholder_body or has_bu: return lvl
    return None

def _bu_kind(ppr):
    """'none', 'char' (a bullet), 'auto' (a number), or None when this level doesn't say."""
    if ppr is None: return None
    if ppr.find(q('a', 'buNone')) is not None: return 'none'
    if ppr.find(q('a', 'buAutoNum')) is not None: return 'auto'
    if ppr.find(q('a', 'buChar')) is not None or ppr.find(q('a', 'buBlip')) is not None: return 'char'
    return None

def _lvl_ppr(lst, lvl):
    return lst.find(q('a', 'lvl%dpPr' % (lvl + 1))) if lst is not None else None

def _ph_match(tree, typ, idx, by_idx):
    """The layout's or master's placeholder that a slide placeholder inherits from."""
    for sp in tree.iter(q('p', 'sp')):
        ph = sp.find('.//' + q('p', 'ph'))
        if ph is None: continue
        if by_idx and idx is not None and ph.get('idx') == idx: return sp
        if not by_idx and ph.get('type', 'body') == typ: return sp
    return None

def bullet_info(ap, sp):
    """(kind, level) of a slide paragraph, following PowerPoint's inheritance: the paragraph, the
    text box's list style, the layout's placeholder, the master's placeholder, the master's text
    styles. kind is 'none', 'char' or 'auto'. A text box that isn't a placeholder has no bullets
    unless it says so."""
    lvl = para_level(ap)
    k = _bu_kind(ap.find(q('a', 'pPr')))
    if k: return k, lvl
    tb = sp.find(q('p', 'txBody')) if sp is not None else None
    k = _bu_kind(_lvl_ppr(tb.find(q('a', 'lstStyle')) if tb is not None else None, lvl))
    if k: return k, lvl
    ph = sp.find('.//' + q('p', 'ph')) if sp is not None else None
    slide = notes.slide if notes is not None else None
    if ph is None or slide is None: return 'none', lvl
    typ, idx = ph.get('type', 'body'), ph.get('idx')
    if typ in ('title', 'ctrTitle', 'sldNum', 'dt', 'ftr', 'hdr'): return 'none', lvl
    master_typ = 'body' if typ in ('obj', 'body') else typ
    layout = slide.slide_layout; master = layout.slide_master
    lay_ph = _ph_match(layout._element, typ, idx, True)
    if lay_ph is None: lay_ph = _ph_match(layout._element, typ, idx, False)
    chain = [lay_ph,
             _ph_match(master._element, master_typ, None, False)]
    for base in chain:
        if base is None: continue
        btb = base.find(q('p', 'txBody'))
        k = _bu_kind(_lvl_ppr(btb.find(q('a', 'lstStyle')) if btb is not None else None, lvl))
        if k: return k, lvl
    styles = master._element.find('.//' + q('p', 'txStyles'))
    body = styles.find(q('p', 'bodyStyle')) if styles is not None else None
    k = _bu_kind(_lvl_ppr(body, lvl))
    return (k or 'none'), lvl

def autonum_of(ap):
    """(type, startAt) for a paragraph PowerPoint numbers automatically, else None."""
    ppr = ap.find(q('a', 'pPr'))
    if ppr is None or ppr.find(q('a', 'buNone')) is not None: return None
    an = ppr.find(q('a', 'buAutoNum'))
    if an is None: return None
    return an.get('type', 'arabicPeriod'), int(an.get('startAt', '1'))

def para_level(ap):
    ppr = ap.find(q('a', 'pPr'))
    return int(ppr.get('lvl', '0')) if ppr is not None else 0

def merge_split_math(items):
    """math, operator-only text, math -> one equation (y_i - y-hat_i split around a text minus)."""
    out = []
    for it in items:
        if (it[0] == 'math' and len(out) >= 2 and out[-1][0] == 't' and out[-2][0] == 'math'
                and ''.join(c for c, _ in out[-1][1]).strip()
                and all(c in OPERATOR_TEXT for c, _ in out[-1][1])):
            ops = ''.join(c for c, _ in out[-1][1]).strip().replace('-', '−')
            out.pop(); base = out[-1][1]
            base.append(mrun(ops))
            for k in list(it[1]): base.append(k)
            stats_inc('split equations merged')
            continue
        out.append(it)
    return out

def superscript_words(sup, before):
    """A superscript number in prose, as words: R-squared, x squared, 10 to the power 4."""
    neg = sup.startswith(('−', '-')); n = sup.lstrip('−-')
    if not neg and n == '2':
        return '-squared' if re.search(r'(?:^|\W)R$', before) else ' squared'
    if not neg and n == '3': return ' cubed'
    return ' to the power ' + ('negative ' if neg else '') + n

def emit_text_paragraph(ap, placeholder_body, tb=None, bold=False, force=False, into=None, num=None, bu=None):
    """Write one slide paragraph; returns the Word paragraph (None if skipped as a repeat).
    into: continue an existing Word paragraph (a fragment joined to the line before it).
    num: (type, number, sequence key) for a numbered step. bu: (kind, level) from bullet_info."""
    key = para_key(ap)
    if key.strip('|') == '': return None
    if bu is None:
        lvl = bullet_level(ap, placeholder_body); kind = 'char' if lvl is not None else 'none'
    else:
        kind, lvl = bu
    style = 'Code' if is_mono(ap, tb) else None
    if into is not None:
        p = into; notes.this_slide.add(key); lead_in = False
        p.add_run(' '); stats_inc('fragments joined')
    else:
        lead_in = bool(notes.prefix) or notes.pending_symbol is not None
        prefix_was = notes.prefix
        p = notes.para(style, key, force=force,
                       redo=lambda: emit_text_paragraph(ap, placeholder_body, tb, bold, force=True, num=num, bu=bu))
        if p is None: return None
        if prefix_was and len(key.replace('|', '').strip()) <= 2:
            review.append(f'{notes.where}: narration "{prefix_was[:60]}" joined to a lone symbol "{key.replace("|", "").strip()}"; check what the arrow points to')
        if num is not None and style is None:
            notes.number(p, num[0], num[1], num[2], lvl or 0)
        elif kind == 'char' and style is None and REAL_BULLETS:
            p.style = notes.doc.styles[BULLET_STYLES[min(lvl or 0, len(BULLET_STYLES) - 1)]]
            stats_inc('bulleted list items')
        elif lvl and style is None:
            p.paragraph_format.left_indent = Pt(18 * min(lvl, 3))
    only_math = all(c.tag in (q('a', 'pPr'), q('a', 'endParaRPr'), q('mc', 'AlternateContent'), q('a14', 'm'))
                    for c in ap) and any(True for _ in ap.iter(q('m', 'oMath')))
    items = []   # ('t', [(ch, rpr)]) | ('math', el)
    chars = []
    def flush():
        if chars: items.append(('t', list(chars))); chars.clear()
    for c in ap:
        c = unwrap(c) if c.tag == q('mc', 'AlternateContent') else c
        if c is None: continue
        if c.tag in (q('a', 'r'), q('a', 'fld')):
            rpr = c.find(q('a', 'rPr'))
            for ch in ''.join(x.text or '' for x in c.iter(q('a', 't'))):
                if ch == ' ': ch = ' '
                if style == 'Code': ch = CODE_QUOTES.get(ch, ch)
                chars.append((ch, rpr))
        elif c.tag == q('a', 'br'):
            # Keynote breaks lines for layout; in prose they become spaces. Code keeps its line
            # breaks.
            chars.append(('\n' if style == 'Code' or not SOFT_BREAKS_AS_SPACES else ' ', None))
        elif c.tag == q('a14', 'm'):
            flush()
            for om in c:
                if om.tag == q('m', 'oMathPara') and not only_math:
                    for inner in om.iter(q('m', 'oMath')): items.append(('math', clean_math(inner)))
                else:
                    items.append(('math', clean_math(om)))
            notes.stats['equations'] += 1
    flush()
    items = merge_split_math(items)
    # short inline equations become words (only inside a sentence; an equation on its own line stays)
    if not only_math and style != 'Code':
        for k, it in enumerate(items):
            if it[0] != 'math': continue
            om = it[1] if local(it[1]) == 'oMath' else next(it[1].iter(q('m', 'oMath')), None)
            if om is not None and is_short_math(om):
                items[k] = ('t', [(ch, None) for ch in math_words(om, notes.where)])
                stats_inc('inline equations to words')
    merged = []
    for it in items:
        if it[0] == 't' and merged and merged[-1][0] == 't': merged[-1] = ('t', merged[-1][1] + it[1])
        else: merged.append(it)
    items = merged
    flat = []
    for idx, it in enumerate(items):
        if it[0] == 't': flat += [(ch, (idx, r)) for ch, r in it[1]]
        else: flat.append(('⁣', (idx, None)))          # placeholder for math: counts as content
    # Code is verbatim: R output's "data:" and quoted strings must survive.
    cleaned = flat if style == 'Code' else clean_chars(flat)
    if style != 'Code':
        out = []                                             # collapse double spaces left by removals
        for c, s in cleaned:
            if c == ' ' and out and out[-1][0] == ' ': continue
            out.append((c, s))
        cleaned = out
    if lead_in and cleaned:
        first = ''.join(ch for ch, _ in cleaned)
        lf = lower_first(first)
        if lf != first:
            cleaned[0] = (lf[0], cleaned[0][1])
    run_text, run_rpr = '', 'unset'
    def add_run(text, rpr):
        if not text: return
        lines = text.split('\n')
        for k, line in enumerate(lines):
            if k: p.add_run().add_break()
            if not line: continue
            base = int(rpr.get('baseline', '0')) if rpr is not None else 0
            if style != 'Code':
                if TIMES_AS_WORD and '×' in line:
                    new_line = re.sub(r'\s*×\s*', ' times ', line)
                    notes.changes.append((notes.where, 'times sign', line.strip(), new_line.strip()))
                    line = new_line; stats_inc('times signs as words')
                if base > 0 and SUPERSCRIPTS_AS_WORDS and re.fullmatch(r'\s*[−-]?\d+\s*', line):
                    before = p.text
                    words = superscript_words(line.strip(), before)
                    sym = (re.findall(r'\S+$', before) or [''])[0]
                    notes.changes.append((notes.where, 'superscript', f'{sym}^{line.strip()}', f'{sym}{words}'))
                    r = p.add_run(words); stats_inc('superscripts as words')
                    if bold: r.bold = True
                    continue
                if base < 0 and SUBSCRIPTS_AS_WORDS and re.fullmatch(r'\s*[\w,]{1,4}\s*', line):
                    runs = p.runs
                    tail = runs[-1].text if runs else ''
                    new_tail = tail
                    if tail and tail[-1] in GREEK_LETTERS:
                        new_tail = tail[:-1] + GREEK_LETTERS[tail[-1]]; runs[-1].text = new_tail
                    word = re.findall(r'\S+$', tail)
                    sym = word[0] if word else ''
                    new_sym = re.findall(r'\S+$', new_tail)[0] if sym else ''
                    notes.changes.append((notes.where, 'subscript', f'{sym}_{line.strip()}', f'{new_sym} sub {line.strip()}'.strip()))
                    r = p.add_run(' sub ' + line.strip()); stats_inc('subscripts as words')
                    if bold: r.bold = True
                    continue
            r = p.add_run(line)
            if bold: r.bold = True
            if base > 0:
                r.font.superscript = True
                review.append(f'{notes.where}: superscript "{line}" in "{key.split("|")[0][:60]}" (a screen reader may not say it; consider words)')
            elif base < 0:
                r.font.subscript = True
    for ch, (idx, rpr) in cleaned:
        if ch == '⁣':
            add_run(run_text, run_rpr if run_rpr != 'unset' else None); run_text, run_rpr = '', 'unset'
            p._p.append(items[idx][1])
            continue
        if run_rpr == 'unset' or rpr is run_rpr:
            run_text += ch; run_rpr = rpr
        else:
            add_run(run_text, run_rpr); run_text, run_rpr = ch, rpr
    add_run(run_text, run_rpr if run_rpr != 'unset' else None)
    return p

LEADS_ON = {'if', 'where', 'when', 'for', 'with', 'and', 'or', 'of', 'to', 'by', 'that', 'the', 'been'}
GREEK_NAMES = {'alpha', 'beta', 'gamma', 'delta', 'epsilon', 'theta', 'lambda', 'mu', 'pi', 'rho',
               'sigma', 'tau', 'phi', 'chi', 'psi', 'omega'}

def continues(a, b, tb, placeholder_body, bu_a=None, bu_b=None):
    """Is paragraph b the rest of a sentence Keynote broke over two lines (same box)?
    Joined when a does not end a sentence and b starts with '(' or a lowercase word that is not a
    symbol being defined ("sigma is", "x = 0", "n, the number", "exp() undoes"); or when a ends on
    "if", "where", ... and b starts with a short equation ("Residuals if" / "beta sub 1 = .1").
    Derived from an instructor's hand edits of generated Word versions."""
    if is_mono(a, tb) or is_mono(b, tb): return False
    if autonum_of(b) is not None: return False                 # a numbered step starts its own line
    if bu_a is not None and bu_b is not None:
        if bu_b[0] != 'none' or bu_a != bu_b: return False     # a bulleted item starts its own line
    elif bullet_level(a, placeholder_body) != bullet_level(b, placeholder_body): return False
    ta = ''.join(x.text or '' for x in a.iter(q('a', 't'))).replace('\u2028', ' ').rstrip()
    if not ta or ta[-1] in '.?!:;…': return False
    last = re.findall(r'[A-Za-z]+', ta)
    for c in b:
        c = unwrap(c) if c.tag == q('mc', 'AlternateContent') else c
        if c is None or c.tag in (q('a', 'pPr'), q('a', 'endParaRPr')): continue
        if c.tag == q('a14', 'm'):
            om = next(c.iter(q('m', 'oMath')), None)
            return bool(last) and last[-1].lower() in LEADS_ON and om is not None and is_short_math(om)
        t = ''.join(x.text or '' for x in c.iter(q('a', 't'))).lstrip()
        if not t: continue
        if t[0] == '(': return True
        if RELATION_START_RE.match(t): return True         # "Population statistic" / "= 42"
        m = re.match(r'([A-Za-z]+)(.{0,3})', t)
        if not m or not t[0].islower(): return False
        word, after = m.group(1), m.group(2)
        if len(word) == 1 or word in GREEK_NAMES: return False
        if after.startswith(('(', ' is', ' =', '=', ',', '_', ' <')): return False
        return True
    return False

def emit_textbody(sp, placeholder_body=False):
    tb = sp.find(q('p', 'txBody'))
    if tb is None: return False
    any_text = False; last_ap = last_p = last_bu = None; first_p = None
    counters = {}        # level -> [type, next number, sequence key]: PowerPoint's own numbering
    paras = [ap for ap in tb.findall(q('a', 'p')) if para_key(ap).strip('|')]
    force_all = False
    depth = {lvl: k for k, lvl in enumerate(sorted({para_level(ap) for ap in paras}))}
    if paras and any(is_mono(ap, tb) for ap in paras):
        # Code repeats as a block, never line by line: a block shown again with one change keeps
        # every line, or its output loses its labels.
        bkey = 'CODE|' + '\n'.join(para_key(ap) for ap in paras)
        if bkey in notes.prev_slide:
            notes.this_slide.add(bkey); notes.stats['repeats skipped'] += len(paras)
            for ap in paras: notes.this_slide.add(para_key(ap))
            return True
        notes.this_slide.add(bkey); force_all = True
    for ap in tb.findall(q('a', 'p')):
        if not para_key(ap).strip('|'): continue
        lvl, an, num = para_level(ap), autonum_of(ap), None
        for L in list(counters):            # a shallower paragraph, or an unnumbered one at the
            if L > lvl or (L == lvl and an is None): del counters[L]   # same level, ends a sequence
        if an is not None:
            c = counters.get(lvl)
            if c is None or c[0] != an[0]:
                notes.num_seq += 1
                c = counters[lvl] = [an[0], an[1], (notes.where, notes.num_seq)]
            num = (c[0], c[1], c[2]); c[1] += 1
        bu = bullet_info(ap, sp)
        bu = (bu[0], depth.get(bu[1], 0))                  # Keynote writes levels 0, 2, 4: nest by rank
        if JOIN_FRAGMENTS and last_p is not None and continues(last_ap, ap, tb, placeholder_body, last_bu, bu):
            emit_text_paragraph(ap, placeholder_body, tb, into=last_p, bu=bu)
            if os.environ.get('PPTX2NOTES_DEBUG'): print('  joined:', last_p.text[-90:])
        else:
            last_p = emit_text_paragraph(ap, placeholder_body, tb, num=num, force=force_all, bu=bu)
            if first_p is None and last_p is not None: first_p = last_p
        last_ap = ap; last_bu = bu; any_text = True
    if first_p is not None:
        b = abs_box(sp)
        if b is not None:
            notes.slide_shapes.append((b, first_p._p)); notes.slide_items.append((b, first_p._p))
    return any_text

def flush_symbol():
    if notes.pending_symbol is not None:
        p = notes.doc.add_paragraph()
        for m in notes.pending_symbol: p._p.append(m)
        notes.pending_symbol = None; notes.pending_box = None

def emit_figure(text, where, force=False):
    flush_symbol()
    if not text:
        review.append(f'{where}: figure with no description')
        p = notes.para(None, None); p.add_run(FIGURE_PREFIX + '[No description]'); return
    p = notes.para(None, 'FIG|' + text, force=force, redo=lambda: emit_figure(text, where, force=True))
    if p is None: return
    p.add_run(FIGURE_PREFIX + clean_text(text))
    notes.stats['figures'] += 1
    marks = sorted(set(re.findall(r'[=%()]', text)))
    if marks:
        review.append(f'{where}: figure description uses {" ".join(marks)}, which a screen reader reads aloud: "{text[:70]}…"')

# ------------------------------------------------------------------ equation groups
def leaves(g):
    for child in g:
        el = unwrap(child)
        if el is None: continue
        tag = local(el)
        if tag in ('nvGrpSpPr', 'grpSpPr', 'extLst'): continue
        if tag == 'grpSp': yield from leaves(el)
        else: yield el

def leaf_text(el):
    return ''.join(x.text or '' for x in el.iter(q('a', 't'))).replace(' ', ' ')

def is_equation_group(g):
    every = list(leaves(g))
    # a drawn part (curve, line, shading), even one marked decorative, makes it a figure
    if any(not leaf_text(l).strip() and not any(True for _ in l.iter(q('m', 'oMath'))) for l in every):
        return False
    ls = [l for l in every if not is_decorative(l)]
    if not ls or not any(True for l in ls for _ in l.iter(q('m', 'oMath'))): return False
    for l in ls:
        if local(l) in ('pic', 'graphicFrame'): return False
        txt = leaf_text(l); has_math = any(True for _ in l.iter(q('m', 'oMath')))
        if not txt.strip() and not has_math: return False      # a drawn shape: a figure, not an equation
        if txt.strip() and not GROUP_TOKEN_RE.match(txt): return False   # words (DV, intercept): a labeled figure
    return True

def leaf_content(l, drop_parens=False, drop_word=None):
    """m: elements for one piece, in its own order: text tokens as math runs, equations unwrapped."""
    out = []
    for ap in l.iter(q('a', 'p')):
        for c in ap:
            c = unwrap(c) if c.tag == q('mc', 'AlternateContent') else c
            if c is None: continue
            if c.tag in (q('a', 'r'), q('a', 'fld')):
                t = ''.join(x.text or '' for x in c.iter(q('a', 't'))).replace(' ', ' ')
                if drop_word: t = re.sub(drop_word, '', t, flags=re.I)
                if drop_parens: t = t.replace('(', '').replace(')', '')
                t = t.strip().replace('-', '−')
                if t: out.append(mrun(t))
            elif c.tag == q('a14', 'm'):
                for om in c.iter(q('m', 'oMath')):
                    for k in clean_math(om): out.append(k)
    return out

def build_group_math(g):
    ls = [l for l in leaves(g) if not is_decorative(l)]
    box = {id(l): (abs_box(l) or [0, 0, 0, 0]) for l in ls}
    order = lambda l: (box[id(l)][0], box[id(l)][1])
    ops = [l for l in ls if re.match(r'\s*(argmin|argmax|min|max)', leaf_text(l), re.I)]
    fors = [l for l in ls if re.match(r'\s*for\b', leaf_text(l), re.I)]
    om = M('oMath')
    if len(ops) == 1 and len(fors) == 1:
        name = re.match(r'\s*(argmin|argmax|min|max)', leaf_text(ops[0]), re.I).group(1).lower()
        rest = sorted([l for l in ls if l is not ops[0] and l is not fors[0]], key=order)
        paren = '(' in leaf_text(ops[0]) and any(')' in leaf_text(l) for l in rest)
        inner = [k for l in rest for k in leaf_content(l, drop_parens=paren)]
        f = M('func'); fn = etree.SubElement(f, q('m', 'fName')); ll = etree.SubElement(fn, q('m', 'limLow'))
        e = etree.SubElement(ll, q('m', 'e')); e.append(mrun(name, plain=True))
        lim = etree.SubElement(ll, q('m', 'lim'))
        for k in leaf_content(fors[0], drop_word=r'^\s*for\s*'): lim.append(k)
        fe = etree.SubElement(f, q('m', 'e'))
        if paren:
            d = etree.SubElement(fe, q('m', 'd')); de = etree.SubElement(d, q('m', 'e'))
            for k in inner: de.append(k)
        else:
            for k in inner: fe.append(k)
        om.append(f)
    else:
        for l in sorted(ls, key=order):
            for k in leaf_content(l): om.append(k)
    return om

def emit_equation_group(g, where, d):
    flush_symbol()
    om = build_group_math(g)
    p = notes.para(None, 'EQG|' + math_text(om))
    if p is None: return
    mp = M('oMathPara'); mp.append(om); p._p.append(mp)
    notes.stats['equations'] += 1; stats_inc('equation groups as live equations')
    review.append(f'{where}: equation group written as one live equation from its pieces; check it against its reading: "{d[:90]}"')

# ------------------------------------------------------------------ equation-shaped figures
WORDS_RE = re.compile(r'([A-Za-z]{2,}(?:\s+[A-Za-z]{2,})*)')

def mtext(text):
    """Upright words inside an equation, spaces kept (Word's normal-text math run)."""
    r = M('r'); rp = etree.SubElement(r, q('m', 'rPr')); etree.SubElement(rp, q('m', 'nor'))
    t = etree.SubElement(r, q('m', 't')); t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return r

def text_math(t):
    """Slide text inside an equation: words (two letters or more) as upright text, the rest as
    math with layout spaces dropped."""
    out = []
    for k, seg in enumerate(WORDS_RE.split(t)):
        if k % 2: out.append(mtext(' '.join(seg.split())))
        else:
            seg = ''.join(seg.split()).replace('-', '−')
            if seg: out.append(mrun(seg))
    return out

def piece_runs(l):
    """('t', text) / ('m', oMath) for one piece, in order."""
    out = []
    for ap in l.iter(q('a', 'p')):
        for c in ap:
            c = unwrap(c) if c.tag == q('mc', 'AlternateContent') else c
            if c is None: continue
            if c.tag in (q('a', 'r'), q('a', 'fld')):
                t = ''.join(x.text or '' for x in c.iter(q('a', 't'))).replace('\u00a0', ' ')
                if out and out[-1][0] == 't': out[-1] = ('t', out[-1][1] + t)   # ".", "0", "3" -> ".03"
                else: out.append(('t', t))
            elif c.tag == q('a14', 'm'):
                for om in c.iter(q('m', 'oMath')): out.append(('m', om))
    return out

def runs_math(runs):
    out = []
    for kind, v in runs:
        if kind == 't': out += text_math(v)
        else: out += list(clean_math(v))
    return out

def has_omath(l): return any(True for _ in l.iter(q('m', 'oMath')))

def is_fraction_bar(l):
    if local(l) not in ('sp', 'cxnSp') or leaf_text(l).strip() or has_omath(l): return False
    g = l.find('.//' + q('a', 'prstGeom')); b = abs_box(l)
    return (g is not None and g.get('prst') in ('line', 'straightConnector1') and b is not None
            and b[2] > 0 and b[3] <= 2 * EMU_PT)

def hand_fraction(g):
    """A fraction built by hand (Z = mu bar minus mu, a drawn line, and a term below it):
    one horizontal line, every other piece an equation or a
    short term, left of it, right of it, above it or below it. Returns the pieces or None."""
    ls = list(leaves(g))
    bars = [l for l in ls if is_fraction_bar(l)]
    if len(bars) != 1: return None
    bar = bars[0]; bx, by, bw, bh = abs_box(bar); tol = 6 * EMU_PT
    left, num, den, right = [], [], [], []
    for l in ls:
        if l is bar: continue
        txt = leaf_text(l).strip()
        if local(l) in ('pic', 'graphicFrame') or is_decorative(l): return None
        if not txt and not has_omath(l): return None           # another drawn part: a figure
        if len(txt.split()) > 4: return None                   # a sentence, not a term
        b = abs_box(l)
        if b is None: return None
        x, y, w, h = b
        if x + w <= bx + tol: left.append(l)
        elif x >= bx + bw - tol: right.append(l)
        elif y + h <= by + tol: num.append(l)
        elif y >= by - tol: den.append(l)
        else: return None
    if not num or not den or not any(has_omath(l) for l in left + num + den + right): return None
    return left, num, den, right

def build_fraction(parts):
    left, num, den, right = parts
    byx = lambda ls: sorted(ls, key=lambda l: (abs_box(l)[0], abs_box(l)[1]))
    om = M('oMath')
    for l in byx(left): om.extend(runs_math(piece_runs(l)))
    f = M('f'); fn = etree.SubElement(f, q('m', 'num')); fd = etree.SubElement(f, q('m', 'den'))
    for l in byx(num): fn.extend(runs_math(piece_runs(l)))
    for l in byx(den): fd.extend(runs_math(piece_runs(l)))
    om.append(f)
    for l in byx(right): om.extend(runs_math(piece_runs(l)))
    return [om]

def icon_name(l):
    cnv = l.find('.//' + q('p', 'cNvPr'))
    d = descr(l) if not is_decorative(l) else ''
    name = d if d and len(d.split()) <= 2 else (cnv.get('name') or '').strip() if cnv is not None else ''
    if not name or GENERIC_NAME_RE.match(name) or not re.match(r'^[A-Za-z][A-Za-z ]{1,30}$', name): return None
    return name.lower() if name[0].isupper() and name[1:].islower() else name

def equation_rows(g):
    """Rows of equations, each with a gap holding a named icon (P( [icon] in s) = 15% = .15,
    where the icon is a picture named for what it shows). Every text piece must be an equation (a relation sign, no words); every other
    piece a picture whose name is a word. Returns one oMath per row, top to bottom, or None."""
    rows, icons = [], []
    for l in leaves(g):
        txt = leaf_text(l)
        if txt.strip() or has_omath(l):
            if is_decorative(l): return None
            if re.search(r'[A-Za-z]{3,}', txt) or not RELATION_RE.search(txt + math_text(l)): return None
            rows.append(l)
        elif local(l) in ('sp', 'pic') and l.find('.//' + q('a', 'blipFill')) is not None:
            icons.append(l)
        else:
            return None
    if len(rows) < 2: return None
    top = {id(r): abs_box(r) for r in rows}
    if any(b is None for b in top.values()): return None
    owner = {}
    for ic in icons:
        b = abs_box(ic); name = icon_name(ic)
        if b is None or name is None: return None
        cy = b[1] + b[3] / 2
        above = [r for r in rows if top[id(r)][1] <= cy]
        if not above: return None
        r = max(above, key=lambda r: top[id(r)][1])
        if id(r) in owner: return None
        owner[id(r)] = name
    out = []
    for r in sorted(rows, key=lambda r: (top[id(r)][1], top[id(r)][0])):
        runs = piece_runs(r); name = owner.get(id(r))
        if name is not None:
            for k, (kind, v) in enumerate(runs):
                m = re.search(r' {3,}', v) if kind == 't' else None
                if m:
                    runs[k:k + 1] = [('t', v[:m.start()]), ('name', name), ('t', v[m.end():])]
                    break
            else:
                return None                                     # an icon with no gap to fill
        elif any(kind == 't' and re.search(r' {3,}', v.strip()) for kind, v in runs):
            return None                                         # a gap with nothing to fill it
        om = M('oMath')
        for kind, v in runs:
            if kind == 'name': om.append(mtext(v))
            elif kind == 't': om.extend(text_math(v))
            else: om.extend(list(clean_math(v)))
        out.append(om)
    return out

def emit_equation_figure(g, where, d):
    """An equation-shaped figure as live equations. True if handled."""
    parts = hand_fraction(g)
    oms = build_fraction(parts) if parts else equation_rows(g)
    if not oms: return False
    if d and all('EQF|' + math_text(om) in notes.prev_slide for om in oms):
        return False      # a repeat whose description says what changed (66: the denominator boxed in red)
    flush_symbol()
    for om in oms:
        p = notes.para(None, 'EQF|' + math_text(om))
        if p is None: continue
        mp = M('oMathPara'); mp.append(om); p._p.append(mp)
        notes.stats['equations'] += 1
    stats_inc('equation-shaped figures as live equations')
    review.append(f'{where}: figure written as {len(oms)} live equation(s) instead of its description; check against it: "{d[:90]}"')
    return True

# ------------------------------------------------------------------ tables
def emit_table(gf, where):
    flush_symbol()
    tbl = gf.find('.//' + q('a', 'tbl'))
    rows = tbl.findall(q('a', 'tr'))
    if not rows: return
    data = [[' '.join(''.join(t.text or '' for t in tc.iter(q('a', 't'))).split()) for tc in tr.findall(q('a', 'tc'))] for tr in rows]
    key = 'TBL|' + repr(data)
    if key in notes.prev_slide:
        notes.stats['repeats skipped'] += 1; notes.this_slide.add(key); return
    notes.this_slide.add(key); notes.last_skipped = None
    for old_data, old_el in list(notes.prev_tables):
        same_shape = len(old_data) == len(data) and all(len(a) == len(b) for a, b in zip(old_data, data))
        if same_shape and all(o == '' or o == n for ro, rn in zip(old_data, data) for o, n in zip(ro, rn)):
            spacer = old_el.getnext()
            old_el.getparent().remove(old_el)
            if spacer is not None and spacer.tag == qn('w:p') and not ''.join(spacer.itertext()).strip():
                spacer.getparent().remove(spacer)
            notes.stats['tables'] -= 1; stats_inc('build tables merged')
            notes.prev_tables.remove((old_data, old_el))
    ncol = max(len(r) for r in data)
    t = notes.doc.add_table(rows=len(data), cols=ncol); t.style = 'Table Grid'
    for i, tr in enumerate(rows):
        for j, tc in enumerate(tr.findall(q('a', 'tc'))):
            par = t.cell(i, j).paragraphs[0]
            for ar in tc.iter(q('a', 'r')):
                txt = ''.join(x.text or '' for x in ar.iter(q('a', 't')))
                if not txt: continue
                run = par.add_run(clean_text(txt) if txt.strip() else txt)
                if i == 0: run.bold = True
                rpr = ar.find(q('a', 'rPr'))
                base = int(rpr.get('baseline', '0')) if rpr is not None else 0
                if base > 0:
                    run.font.superscript = True
                    review.append(f'{where}: superscript "{txt}" in a table cell (a screen reader may not say it)')
                elif base < 0:
                    run.font.subscript = True
    trpr = t.rows[0]._tr.get_or_add_trPr()
    h = etree.SubElement(trpr, qn('w:tblHeader')); h.set(qn('w:val'), 'true')
    if notes.prefix:
        review.append(f'{where}: narration "{notes.prefix}" led into a table; placed before it')
    notes.doc.add_paragraph()
    notes.stats['tables'] += 1
    notes.this_slide_tables.append((data, t._tbl))
    b = abs_box(gf)
    if b is not None:
        notes.slide_tables.append((b, [int(tr.get('h', '0')) for tr in rows], t._tbl))
        notes.slide_items.append((b, None))

# ------------------------------------------------------------------ labels, narration, annotations
def word_stems(text):
    return {w[:-1] if len(w) > 3 and w.endswith('s') else w for w in re.findall(r'[a-z]+', text.lower())}

def box_gap(a, b):
    """Distance between two boxes (0 if they overlap)."""
    dx = max(b[0] - (a[0] + a[2]), a[0] - (b[0] + b[2]), 0)
    dy = max(b[1] - (a[1] + a[3]), a[1] - (b[1] + b[3]), 0)
    return max(dx, dy)

def collect_figures(slide):
    """Described figures on the slide, for labels laid over them."""
    for el in slide.shapes._spTree.iter(q('p', 'grpSp'), q('p', 'pic')):
        d = descr(el)
        if d and not is_decorative(el):
            b = abs_box(el)
            if b is not None: notes.fig_boxes.append((b, word_stems(d)))

def covered_label(el, paras, maths, text_only):
    """A short label inside a described figure whose words its description already says (two
    Tail labels over a chart described as having arrows pointing to the tails). It should have
    been grouped with the figure; the description covers it either way."""
    words = re.findall(r'[A-Za-z]+', text_only)
    if maths or len(paras) != 1 or not 1 <= len(words) <= 3: return False
    b = abs_box(el)
    if b is None: return False
    cx, cy = b[0] + b[2] / 2, b[1] + b[3] / 2
    for fb, stems in notes.fig_boxes:
        if fb[0] <= cx <= fb[0] + fb[2] and fb[1] <= cy <= fb[1] + fb[3] and word_stems(text_only) <= stems:
            return True
    return False

def is_lead_label(el):
    if local(el) != 'sp' or (ph_info(el) or ('', ''))[0] in ('title', 'ctrTitle'): return False
    paras = [ap for ap in el.iter(q('a', 'p')) if para_key(ap).strip('|')]
    if len(paras) != 1 or any(True for _ in el.iter(q('m', 'oMath'))): return False
    t = ''.join(x.text or '' for x in paras[0].iter(q('a', 't'))).strip()
    return t.endswith(':') and 1 <= len(t.split()) <= 5

def reorder_lead_labels(children):
    """A label ending in a colon just above a text box, but read after it (Two-tailed test:
    above An observation in the left or right tail ...), is moved before it and leads into it.
    Header before body; the deck's own order is still wrong
    for the PDF."""
    els = [unwrap(c) for c in children]
    order = list(range(len(els)))
    for i, el in enumerate(els):
        if el is None or is_decorative(el) or not is_lead_label(el): continue
        lb = abs_box(el)
        if lb is None: continue
        best = None
        for j in range(i):
            t = els[j]
            if t is None or local(t) != 'sp' or is_decorative(t) or not any(para_key(ap).strip('|') for ap in t.iter(q('a', 'p'))): continue
            tb = abs_box(t)
            if tb is None: continue
            gap = tb[1] - (lb[1] + lb[3])
            overlap = min(lb[0] + lb[2], tb[0] + tb[2]) - max(lb[0], tb[0])
            if -5 * EMU_PT <= gap <= 80 * EMU_PT and overlap >= 0.5 * lb[2]:
                if best is None or gap < best[0]: best = (gap, j)
        if best is not None:
            j = best[1]; order.remove(i); order.insert(order.index(j), i)
            notes.lead_labels.add(id(el))
    return [children[k] for k in order]

def is_narration_group(g):
    """A group of bare arrows or lines carrying alt text (two arrows from two steps of a list,
    The comparison and decision is:) is narration, like a single narrated arrow."""
    ls = list(leaves(g))
    if not ls: return False
    for l in ls:
        if local(l) not in ('sp', 'cxnSp') or leaf_text(l).strip() or has_omath(l): return False
        if l.find('.//' + q('a', 'blipFill')) is not None: return False
        geom = l.find('.//' + q('a', 'prstGeom'))
        if geom is None or geom.get('prst') not in NARRATION_GEOMS: return False
    return True

def line_ends(el):
    """Both ends of a narrated line, or of every line in a narrated group."""
    pts = []
    for l in ([el] if local(el) != 'grpSp' else list(leaves(el))):
        b = abs_box(l); sppr = l.find(q('p', 'spPr'))
        x = sppr.find(q('a', 'xfrm')) if sppr is not None else None
        if b is None or x is None: continue
        fh, fv = x.get('flipH') == '1', x.get('flipV') == '1'
        pts.append((b[0] + (b[2] if fh else 0), b[1] + (b[3] if fv else 0)))
        pts.append((b[0] + (0 if fh else b[2]), b[1] + (0 if fv else b[3])))
    return pts

def point_gap(pt, b):
    dx = max(b[0] - pt[0], 0, pt[0] - (b[0] + b[2])); dy = max(b[1] - pt[1], 0, pt[1] - (b[1] + b[3]))
    return (dx * dx + dy * dy) ** 0.5

def place_trailing_narration(where):
    """Narration with nothing after it on the slide goes before the item nearest its arrow (The
    comparison and decision is / (This is where p-values come in))."""
    if not notes.prefix or notes.prefix_box is None or not notes.slide_items: return False
    ends = notes.prefix_points
    if ends:
        # The narration joins the items at its ends; it goes before the one read later, where a
        # narrated arrow belongs in reading order. A table can't take it.
        near = [p for b, p in notes.slide_items if min(point_gap(e, b) for e in ends) <= 40 * EMU_PT]
        if not near: return False
        p = near[-1]
        if p is None or p in notes.prefixed: return False
    else:
        if not notes.slide_shapes: return False
        gap, p = min(((box_gap(notes.prefix_box, b), p) for b, p in notes.slide_shapes), key=lambda x: x[0])
        if gap > 60 * EMU_PT or p in notes.prefixed: return False
    runs = p.findall(qn('w:r'))
    first_t = next((t for r in runs for t in r.iter(qn('w:t')) if t.text), None)
    if first_t is not None:
        lead = '(' if first_t.text.startswith('(') else ''
        first_t.text = lead + lower_first(first_t.text[len(lead):])
    r = etree.Element(qn('w:r')); t = etree.SubElement(r, qn('w:t')); t.text = notes.prefix + ' '
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    ppr = p.find(qn('w:pPr'))
    if ppr is not None: ppr.addnext(r)
    else: p.insert(0, r)
    review.append(f'{where}: narration "{notes.prefix[:60]}" had nothing after it; placed before the nearest item, "{"".join(x.text or "" for x in p.iter(qn("w:t")))[:60]}"')
    stats_inc('trailing narration placed'); notes.prefix = ''; notes.prefix_box = None
    return True

def annotate_table_row(el, text_only):
    """A wordless annotation beside a table row (< 0.01 to the right of a row)
    goes at the end of that row's last cell, where it sits on the slide."""
    b = abs_box(el)
    if b is None: return False
    cy = b[1] + b[3] / 2
    for tb, heights, tbl in notes.slide_tables:
        right = tb[0] + tb[2]
        if not (right - 10 * EMU_PT <= b[0] <= right + 80 * EMU_PT): continue
        y = tb[1]
        for k, h in enumerate(heights):
            if y <= cy < y + h:
                trs = tbl.findall(qn('w:tr'))
                if k >= len(trs) or k == 0: return False
                cell_p = trs[k].findall(qn('w:tc'))[-1].findall(qn('w:p'))[-1]
                r = etree.SubElement(cell_p, qn('w:r')); t = etree.SubElement(r, qn('w:t'))
                t.text = ' ' + clean_text(text_only)
                t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                row_text = ''.join(x.text or '' for x in trs[k].iter(qn('w:t')))
                review.append(f'{notes.where}: annotation "{text_only}" beside the table added to its row: "{row_text[:60]}"')
                stats_inc('table annotations into rows')
                return True
            y += h
    return False

# ------------------------------------------------------------------ walk
def walk(container, where, title_el):
    children = list(container)
    if container.tag == q('p', 'spTree'): children = reorder_lead_labels(children)
    for child in children:
        el = unwrap(child)
        if el is None: continue
        tag = local(el)
        if tag in ('nvGrpSpPr', 'grpSpPr', 'extLst'): continue
        if el is title_el: continue
        if is_decorative(el): continue
        info = ph_info(el) if tag == 'sp' else None
        if info and info[0] in SKIP_PH: continue
        d = descr(el)
        if tag == 'grpSp':
            if d and is_narration_group(el):
                if notes.last_skipped is not None:
                    redo = notes.last_skipped; notes.last_skipped = None; redo()
                    stats_inc('repeats kept before narration')
                notes.prefix = (notes.prefix + ' ' + clean_text(d)).strip(); notes.prefix_box = abs_box(el)
                notes.prefix_points = line_ends(el)
                stats_inc('narration groups')
            elif is_equation_group(el): emit_equation_group(el, where, d)
            elif emit_equation_figure(el, where, d): pass
            elif d: emit_figure(d, where)
            else:
                ls = list(leaves(el))
                if any(local(x) in ('graphicFrame', 'pic') for x in ls):
                    for l in ls:
                        t = leaf_text(l).strip()
                        if t and len(t.split()) <= 2 and not has_omath(l) and not is_decorative(l):
                            review.append(f'{where}: label "{t}" comes from a group with no description (a table or picture with loose labels); describe the group, or mark the label decorative')
                walk(el, where, title_el)
        elif tag == 'pic':
            emit_figure(d, where)
        elif tag == 'graphicFrame':
            if el.find('.//' + q('a', 'tbl')) is not None: emit_table(el, where)
            elif d: emit_figure(d, where)
        elif tag in ('sp', 'cxnSp'):
            has_text = tag == 'sp' and any(para_key(ap).strip('|') for ap in el.iter(q('a', 'p')))
            if has_text:
                paras = [ap for ap in el.iter(q('a', 'p')) if para_key(ap).strip('|')]
                maths = [m for ap in paras for m in ap.iter(q('m', 'oMath'))]
                text_only = ''.join(x.text or '' for ap in paras for x in ap.iter(q('a', 't'))).strip()
                sym = ''.join(math_text(m) for m in maths).replace(':', '').strip()
                if covered_label(el, paras, maths, text_only):
                    stats_inc('labels left to the figure description'); continue
                if id(el) in notes.lead_labels:
                    notes.prefix = (notes.prefix + ' ' + clean_text(text_only.rstrip().rstrip(':')) + ',').strip()
                    notes.prefix_box = None
                    stats_inc('labels read before the text below them')
                    review.append(f'{where}: label "{text_only}" read after the text below it; written as its lead-in (fix the reading order in the deck for the PDF)')
                    continue
                if (not maths and len(paras) == 1 and RELATION_START_RE.match(text_only)
                        and not re.search(r'[A-Za-z]{2,}', text_only)):
                    if annotate_table_row(el, text_only): continue
                    review.append(f'{where}: "{text_only}" stands alone (an annotation with no words); say what it refers to in the deck')
                if maths and not text_only and 0 < len(sym) <= 4 and not RELATION_RE.search(sym):
                    flush_symbol()
                    ms = [clean_math(m) for m in maths]
                    strip_trailing_colon(ms[-1])
                    notes.pending_symbol = ms; notes.pending_box = abs_box(el)
                    continue
                if notes.pending_symbol is not None:
                    tb = el.find(q('p', 'txBody'))
                    words = re.search(r'[A-Za-z]{2,}', text_only)
                    code = bool(paras) and is_mono(paras[0], tb)
                    if not words or code or not beside(notes.pending_box, abs_box(el)):
                        flush_symbol()
                emit_textbody(el, placeholder_body=bool(info and info[0] in ('body', 'obj')))
            elif d:
                if notes.last_skipped is not None:      # the narration continues from a repeated item: keep it
                    redo = notes.last_skipped; notes.last_skipped = None; redo()
                    stats_inc('repeats kept before narration')
                notes.prefix = (notes.prefix + ' ' + clean_text(d)).strip()   # narrated arrow / labeled shape
                notes.prefix_box = abs_box(el); notes.prefix_points = line_ends(el)

def agenda_items(slide):
    items = []
    for sh in slide.shapes:
        if not sh.has_text_frame or sh == slide.shapes.title: continue
        for para in sh.text_frame.paragraphs:
            t = ''.join(r.text for r in para.runs).strip()
            if t: items.append((para.level, t))
    return items

def norm_title(t): return re.sub(r'[^a-z0-9]+', ' ', t.lower()).strip()

def is_title_slide(s): return s.slide_layout.name.lower().startswith('title')

def title_slide_lines(s):
    """[[line, ...] per text shape] on the title slide, in reading order, decorative skipped."""
    out = []
    for sh in s.shapes:
        if not sh.has_text_frame or is_decorative(sh._element): continue
        lines = [' '.join(p.text.split()) for p in sh.text_frame.paragraphs]
        lines = [l for l in lines if l]
        if lines: out.append(lines)
    return out

def front_matter(s, title):
    """Course name, instructor and pronouns from the title slide, under the document title.
    A line the title already says is left out."""
    nt = norm_title(title)
    for lines in title_slide_lines(s):
        keep = [clean_text(l) for l in lines if norm_title(l) and norm_title(l) not in nt]
        if keep:
            notes.doc.add_paragraph(', '.join(keep)); stats_inc('title page lines')

# ------------------------------------------------------------------ slide headings
def slide_heading(first, last, title):
    t = clean_title(title)
    fmt = BUILD_HEADING_FMT if last != first else SLIDE_HEADING_FMT
    h = fmt.format(n=first, first=first, last=last, title=t)
    if not t: h = re.sub(r'[\s:.,;-]+$', '', h)          # "Slide 12" for a slide with no title
    return h

# ------------------------------------------------------------------ R output as tables
R_STARS = re.compile(r'^[*.]+$')
R_NAME = re.compile(r'^(?:\(Intercept\)|[A-Za-z.][\w.]*)$')
R_NUM = re.compile(r'^[<>]?-?(?:\d+\.?\d*|\.\d+)(?:e[-+]?\d+)?$|^NA$', re.I)
R_KINDS = [
    ('coef', re.compile(r'^\s*Estimate\s+Std\.\s*Error\s+([tz])\s+value\s+Pr\(>\|[tz]\|\)\s*$')),
    ('aov', re.compile(r'^\s*Df\s+Sum\s+Sq\s+Mean\s+Sq\s+F\s+value\s+Pr\(>F\)\s*$')),
    ('anova', re.compile(r'^\s*Res\.Df\s+RSS\s+Df\s+Sum\s+of\s+Sq\s+F\s+Pr\(>F\)\s*$')),
    ('lavaan', re.compile(r'^\s*Estimate\s+Std\.Err\s+z-value\s+P\(>\|z\|\)')),
]
R_NOT_ROW = ('Signif', '---', 'Residual standard error', 'F-statistic', 'Multiple R-squared', 'Begin ', 'End ')

def _r_toks(line): return re.sub(r'([<>])\s+', r'\1', line.strip()).split()

def _r_split_stars(tk): return (tk[:-1], tk[-1]) if tk and R_STARS.match(tk[-1]) else (tk, '')

def r_kind(line):
    for kind, rx in R_KINDS:
        m = rx.match(line)
        if m: return kind, m
    return None, None

def is_named_vector(head, values):
    h, v = _r_toks(head), _r_toks(values)
    return (len(h) >= 2 and len(h) == len(v) and all(R_NAME.match(x) for x in h)
            and all(R_NUM.match(x) for x in v) and not head.lstrip().startswith(('#', '>')))

def r_output_tables(doc, code_id):
    """Table-like R output in Code paragraphs becomes a real Word table, after a lead-in line,
    "R output as a table". Returns the number of tables made."""
    from docx.text.paragraph import Paragraph
    W_P = qn('w:p'); body = doc.element.body
    def is_code(e):
        if e is None or e.tag != W_P: return False
        ppr = e.find(qn('w:pPr')); ps = ppr.find(qn('w:pStyle')) if ppr is not None else None
        return ps is not None and ps.get(qn('w:val')) == code_id
    def text(e): return Paragraph(e, doc).text
    def new_p(text_, style_id, before):
        el = etree.Element(W_P); ppr = etree.SubElement(el, qn('w:pPr'))
        if style_id: etree.SubElement(ppr, qn('w:pStyle')).set(qn('w:val'), style_id)
        r = etree.SubElement(el, qn('w:r')); t = etree.SubElement(r, qn('w:t')); t.text = text_
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        before.addprevious(el); return el
    # a Code paragraph holding several lines of table-like output is split into one per line
    for e in list(body.iterchildren(W_P)):
        if not is_code(e): continue
        lines = text(e).split('\n')
        if len(lines) < 2: continue
        if not any(r_kind(l)[0] for l in lines) and not any(is_named_vector(a, b) for a, b in zip(lines, lines[1:])):
            continue
        for l in lines: new_p(l, code_id, e)
        body.remove(e)
    made = 0
    for e in list(body.iterchildren()):
        if e.getparent() is None or not is_code(e): continue
        h = text(e)
        kind, m = r_kind(h)
        nxt = e.getnext()
        if kind is None and is_code(nxt) and is_named_vector(h, text(nxt)):
            kind = 'vector'
        if kind is None: continue
        lines, cur = [], e.getnext()
        if kind == 'vector':
            lines, cur = [nxt], nxt.getnext()
        else:
            while is_code(cur):
                t = text(cur)
                if not t.strip() or t.strip().startswith(R_NOT_ROW) or r_kind(t)[0]: break
                if not (re.search(r'\d', t) or t.rstrip().endswith('~')): break
                lines.append(cur); cur = cur.getnext()
        if not lines: continue
        rl = [text(x) for x in lines]
        rows, stars = [], []
        if kind == 'coef':
            header = ['Term', 'Estimate', 'Std. Error', f'{m.group(1)} value', 'p-value']
            for l in rl:
                tk, s_ = _r_split_stars(_r_toks(l))
                if len(tk) < 2: continue
                n_vals = min(4, len(tk) - 1)
                rows.append([' '.join(tk[:-n_vals])] + tk[-n_vals:] + [''] * (4 - n_vals)); stars.append(s_)
        elif kind in ('aov', 'anova'):
            header = (['Term', 'Df', 'Sum Sq', 'Mean Sq', 'F value', 'p-value'] if kind == 'aov' else
                      ['Model', 'Res.Df', 'RSS', 'Df', 'Sum of Sq', 'F', 'p-value'])
            for l in rl:
                tk, s_ = _r_split_stars(_r_toks(l))
                rows.append((tk + [''] * len(header))[:len(header)]); stars.append(s_)
        elif kind == 'lavaan':
            cols = ['p-value' if c.startswith('P(>') else c for c in _r_toks(h)]
            paths = any(l.rstrip().endswith('~') for l in rl)
            header = (['Path', 'Label'] if paths else ['Quantity']) + cols
            lhs = ''
            for l in rl:
                tk = _r_toks(l)
                if tk[-1] in ('~', '=~', '~~'): lhs = f'{tk[0]} {tk[-1]}'; continue
                lab = tk[1].strip('()') if len(tk) > 1 and tk[1].startswith('(') else ''
                vals = tk[2:] if lab else tk[1:]
                first = f'{lhs} {tk[0]}' if (paths and lhs) else tk[0]
                rows.append(([first, lab] if paths else [first]) + vals); stars.append('')
            header += [''] * (max(len(r) for r in rows) - len(header))
        else:   # a named vector
            header = _r_toks(h); rows = [_r_toks(rl[0])]; stars = ['']
            review.append(f'R output read as a named vector and made a table: {" ".join(header)[:60]}; check it')
        if any(stars):
            header = header + ['Significance']
            rows = [r + [s_] for r, s_ in zip(rows, stars)]
        width = len(header)
        rows = [(r + [''] * width)[:width] for r in rows]
        # what follows: the R-squared line of a summary() becomes a sentence after the table
        rsq_el = cur if (kind == 'coef' and is_code(cur) and text(cur).strip().startswith('Multiple R-squared')) else None
        anchor = rsq_el.getnext() if rsq_el is not None else cur
        if anchor is None:
            anchor = new_p('', None, body[-1]) if body[-1].tag == qn('w:sectPr') else etree.SubElement(body, W_P)
        prev = e.getprevious()
        label = text(prev).strip() if is_code(prev) else ''
        lead = 'R output as a table'
        if label.endswith(':') and len(label.split()) <= 3:
            if label != 'Coefficients:': lead += ', ' + label.rstrip(':')
            body.remove(prev)
        new_p(lead, None, anchor)
        t = doc.add_table(rows=1 + len(rows), cols=width); t.style = 'Table Grid'
        for j, c in enumerate(header):
            t.rows[0].cells[j].paragraphs[0].add_run(c).bold = True
        for i, r in enumerate(rows, 1):
            for j, v in enumerate(r): t.rows[i].cells[j].paragraphs[0].add_run(v)
        trpr = t.rows[0]._tr.get_or_add_trPr()
        etree.SubElement(trpr, qn('w:tblHeader')).set(qn('w:val'), 'true')
        tp = t._tbl.tblPr; look = tp.find(qn('w:tblLook'))
        if look is None: look = etree.SubElement(tp, qn('w:tblLook'))
        look.set(qn('w:firstRow'), '1'); look.set(qn('w:firstColumn'), '1')
        anchor.addprevious(t._tbl)
        notes.r_tables.add(t._tbl)
        if rsq_el is not None:
            vals = re.findall(r'-?\d*\.\d+|-?\d+', text(rsq_el))
            if len(vals) >= 2:
                new_p(f'Multiple R-squared {vals[0]}, adjusted R-squared {vals[1]}', None, anchor)
                body.remove(rsq_el)
        for x in [e] + lines: body.remove(x)
        made += 1
    return made

# ------------------------------------------------------------------ spacing
def tidy_spacing(doc, code_id):
    """Space after code blocks and tables as paragraph spacing, not empty paragraphs, which a
    screen reader announces as "blank". An empty paragraph stays only between two tables, where
    Word needs one."""
    from docx.text.paragraph import Paragraph
    W_P, W_TBL = qn('w:p'), qn('w:tbl'); body = doc.element.body
    def style(e):
        ppr = e.find(qn('w:pPr')); ps = ppr.find(qn('w:pStyle')) if ppr is not None else None
        return ps.get(qn('w:val')) if ps is not None else None
    def empty(e):
        return (e.tag == W_P and not ''.join(x.text or '' for x in e.iter(qn('w:t'))).strip()
                and not any(True for _ in e.iter(q('m', 'oMath'))) and e.find('.//' + qn('w:drawing')) is None)
    removed = 0
    for e in list(body.iterchildren()):
        if not empty(e): continue
        prv, nxt = e.getprevious(), e.getnext()
        if nxt is None or nxt.tag == qn('w:sectPr'): continue
        if prv is not None and prv.tag == W_TBL and nxt.tag == W_TBL: continue
        body.remove(e); removed += 1
    for e in body.iterchildren(W_P, W_TBL):
        nxt = e.getnext()
        if e.tag == W_P and style(e) == code_id and not (nxt is not None and nxt.tag == W_P and style(nxt) == code_id):
            Paragraph(e, doc).paragraph_format.space_after = Pt(12)
        if e.tag == W_TBL and nxt is not None and nxt.tag == W_P:
            Paragraph(nxt, doc).paragraph_format.space_before = Pt(12)
    return removed

# ------------------------------------------------------------------ large numbers
_ONES = ('zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen '
         'fifteen sixteen seventeen eighteen nineteen').split()
_TENS = ' ten twenty thirty forty fifty sixty seventy eighty ninety'.split(' ')

def _below_1000(n):
    h, r = divmod(n, 100); out = []
    if h: out.append(_ONES[h] + ' hundred')
    if r:
        if r < 20: out.append(_ONES[r])
        else:
            t, o = divmod(r, 10)
            out.append(_TENS[t] + (('-' if NUMBER_WORD_HYPHENS else ' ') + _ONES[o] if o else ''))
    return ' '.join(out)

def number_words(n):
    """Exact whole numbers in words, without "and": 2413 -> two thousand four hundred thirteen."""
    if n == 0: return 'zero'
    out = []
    for value, name in ((10 ** 12, 'trillion'), (10 ** 9, 'billion'), (10 ** 6, 'million'), (1000, 'thousand'), (1, '')):
        k, n = divmod(n, value)
        if k: out.append((_below_1000(k) if k < 1000 else number_words(k)) + (' ' + name if name else ''))
    return ' '.join(out)

def number_said(whole, frac, dollar, in_table):
    """The number rule: None when the number is left as it is."""
    if whole < 1000: return None
    if frac or in_table:
        s = f'{whole}{frac}'                                   # digits without the comma
    elif whole % 1000 == 0 and whole < 10 ** 6:
        k = whole // 1000                                      # 140 thousand; one thousand
        s = f'{number_words(k) if k < 10 else k} thousand'
    elif whole >= 10 ** 9 and whole % 10 ** 8 == 0:
        s = f'{whole / 10 ** 9:g} billion'
    elif whole >= 10 ** 6 and whole % 10 ** 5 == 0:
        s = f'{whole / 10 ** 6:g} million'
    else:
        s = number_words(whole)
    return s + (' dollars' if dollar else '')

NUM_RE = re.compile(r'(?<![\w.,$])(\$?)(\d{1,3}(?:,\d{3})+|\d{4,})(\.\d+)?(?![\w$]|,\d|\.\d)')

def number_pass(doc, code_id):
    """Numbers of 1,000 or more outside code, equations, and R output tables. Records every change."""
    from docx.text.paragraph import Paragraph
    W_P, W_TBL = qn('w:p'), qn('w:tbl'); body = doc.element.body
    where = ['the start']
    def style(e):
        ppr = e.find(qn('w:pPr')); ps = ppr.find(qn('w:pStyle')) if ppr is not None else None
        return ps.get(qn('w:val')) if ps is not None else ''
    def fix_para(e, in_table):
        for t in e.iter(qn('w:t')):
            old = t.text or ''
            def rep(m):
                dollar, digits, frac = m.group(1), m.group(2), m.group(3) or ''
                whole = int(digits.replace(',', ''))
                if re.search(r'\b(?:pp?\.|pages?)\s*(?:\d+\s*[–-]\s*)?$', old[:m.start()], re.I):
                    stats_inc('page numbers left as digits'); return m.group(0)
                if not dollar and not frac and ',' not in digits and 1600 <= whole <= 2099:
                    stats_inc('possible years left as digits')
                    notes.changes.append((where[0], 'left as a year', m.group(0), m.group(0)))
                    return m.group(0)
                if in_table and ',' not in digits: return m.group(0)
                said = number_said(whole, frac, bool(dollar), in_table)
                if said is None or said == m.group(0): return m.group(0)
                notes.changes.append((where[0], 'number', m.group(0), said)); stats_inc('numbers rewritten')
                return said
            new = NUM_RE.sub(rep, old)
            if new != old: t.text = new
    for e in body.iterchildren(W_P, W_TBL):
        if e.tag == W_P:
            st = style(e)
            txt = ''.join(x.text or '' for x in e.iter(qn('w:t')))
            m = re.match(r'\s*Slides?\s+(\d+)(?:\s+to\s+(\d+))?', txt)
            if st.lower().startswith('heading') and m:
                where[0] = f'slide {m.group(1)}' + (f' to {m.group(2)}' if m.group(2) else '')
            if st == code_id: continue
            fix_para(e, False)
        elif e not in notes.r_tables:
            for p in e.iter(W_P): fix_para(p, True)

# ------------------------------------------------------------------ code markers
def add_code_markers(doc, code_id):
    """Begin before each run of Code paragraphs; End after it when plain text follows
    (not before a figure, an equation, a heading or a table)."""
    W_P, W_TBL = qn('w:p'), qn('w:tbl')
    def style_of(e):
        ppr = e.find(qn('w:pPr')); ps = ppr.find(qn('w:pStyle')) if ppr is not None else None
        return ps.get(qn('w:val')) if ps is not None else 'Normal'
    def text_of(e): return ''.join(x.text or '' for x in e.iter(qn('w:t')))
    def has_math(e): return any(True for _ in e.iter(q('m', 'oMath')))
    def is_code(e): return e.tag == W_P and style_of(e) == code_id
    def empty(e): return e.tag == W_P and not text_of(e).strip() and not has_math(e)
    def needs_end(e):
        if e is None or e.tag == W_TBL: return False
        st = style_of(e)
        if st.lower().startswith(('heading', 'title')): return False
        t = text_of(e).strip()
        if t.startswith(FIGURE_PREFIX.strip()): return False
        if not t and has_math(e): return False                 # an equation
        return bool(t)
    def marker(text):
        p = etree.Element(W_P); ppr = etree.SubElement(p, qn('w:pPr'))
        ps = etree.SubElement(ppr, qn('w:pStyle')); ps.set(qn('w:val'), code_id)
        r = etree.SubElement(p, qn('w:r')); t = etree.SubElement(r, qn('w:t')); t.text = text
        return p
    els = [e for e in doc.element.body if e.tag in (W_P, W_TBL)]
    i = nb = ne = 0
    while i < len(els):
        if not is_code(els[i]): i += 1; continue
        j = i
        while j + 1 < len(els) and (is_code(els[j + 1]) or (empty(els[j + 1]) and j + 2 < len(els) and is_code(els[j + 2]))):
            j += 1
        prev = els[i].getprevious()
        while prev is not None and prev.tag == W_P and empty(prev): prev = prev.getprevious()
        if prev is not None and prev.tag == W_P and IN_R_RE.match(text_of(prev)):
            prev.getparent().remove(prev)                    # the slide's own "In R:" label: the marker replaces it
        els[i].addprevious(marker(f'Begin {CODE_LANGUAGE} code block.')); nb += 1
        k = j + 1
        while k < len(els) and empty(els[k]): k += 1
        if needs_end(els[k] if k < len(els) else None):
            els[j].addnext(marker(f'End {CODE_LANGUAGE} code block.')); ne += 1
        i = j + 1
    return nb, ne

# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser(description='Turn a remediated .pptx into a text-only companion document (.docx).')
    ap.add_argument('deck', help='the remediated .pptx')
    ap.add_argument('-o', '--out', help='write only this .docx (no copy in _snapshots/)')
    ap.add_argument('--changes', help='also write every number and symbol rewritten to this CSV file')
    args = ap.parse_args()
    prs = Presentation(args.deck)
    base = os.path.splitext(args.deck)[0]
    all_slides = list(prs.slides)
    hidden = [i + 1 for i, s in enumerate(all_slides) if s._element.get('show') == '0']
    if hidden:
        review.append(f'hidden slides skipped and not numbered, as in the exported PDF: {", ".join(map(str, hidden))}')
    slides = [s for s in all_slides if s._element.get('show') != '0']
    ts = slides[0] if slides and is_title_slide(slides[0]) else None
    title = prs.core_properties.title or ''
    if not title and ts is not None:            # no Title property: the title slide's "Week N: ..." line
        title = next((l for ls in title_slide_lines(ts) for l in ls if re.match(r'week\s*\d', l, re.I)), '')
        if title: review.append(f'no Title property in the deck; document title taken from the title slide: "{title}"')
    title = title or os.path.basename(base)
    global notes; notes = Notes(title)
    if ts is not None: front_matter(ts, title)
    agendas = [i for i, s in enumerate(slides) if s.shapes.title is not None and AGENDA_RE.match(s.shapes.title.text_frame.text)]
    sub_items = []
    if len(agendas) > 1:
        lv = agenda_items(slides[agendas[1]])
        deepest = max(l for l, _ in lv)
        sub_items = [t for l, t in lv if l == deepest]
    section_no = 0
    title_level = 3 if len(agendas) > 1 else 2  # slide titles sit under the section headings
    run = None                                  # the heading the last slide wrote: a build can join it
    for i, s in enumerate(slides):
        n = i + 1                               # the slide's page in the exported PDF
        where = f'slide {n}'
        notes.new_slide(where); notes.slide = s
        if i == 0 and ts is not None:
            run = None; continue               # title slide: the document title and front matter
        t_el = s.shapes.title._element if s.shapes.title is not None else None
        tt = ' '.join(s.shapes.title.text_frame.text.replace(' ', ' ').split()) if t_el is not None else ''
        agenda = None                          # 'first', 'section' or 'closing'
        if i in agendas:
            later = [j for j in range(i + 1, len(slides)) if j not in agendas]
            agenda = 'closing' if not later else 'first' if agendas.index(i) == 0 else 'section'
        if not SLIDE_HEADINGS:
            # the earlier format: agendas as before, titles as headings only when they change
            if agenda == 'closing': continue
            if agenda == 'first':
                notes.doc.add_heading(clean_title(tt) or "Today's agenda", level=2)
                for lvl, t in agenda_items(s):
                    ap_ = notes.doc.add_paragraph(clean_text(t))
                    if lvl: ap_.paragraph_format.left_indent = Pt(9 * min(lvl, 4))
                notes.last_title = None; continue
            if agenda == 'section':
                if section_no < len(sub_items):
                    h = clean_text(sub_items[section_no])
                    notes.doc.add_heading(h, level=2); notes.section_heading = h
                else:
                    review.append(f'{where}: agenda slide with no matching section name')
                section_no += 1; notes.last_title = None; continue
            if t_el is not None and not is_decorative(t_el) and tt and tt != notes.last_title:
                emit = True
                if notes.section_heading:
                    a, b = norm_title(tt), norm_title(notes.section_heading)
                    if a == b:
                        emit = False; stats_inc('titles same as section heading dropped')
                    elif a in b or b in a or difflib.SequenceMatcher(None, a, b).ratio() >= 0.8:
                        review.append(f'{where}: slide title "{tt}" closely matches the section heading "{notes.section_heading}"; keep both?')
                    notes.section_heading = None
                if emit:
                    notes.doc.add_heading(clean_title(tt), level=title_level); stats_inc('slide-title headings')
                notes.last_title = tt
        else:
            # every slide gets a numbered heading, its title as on the slide (repeats included)
            level = title_level
            if agenda == 'section':
                if section_no < len(sub_items):
                    h = clean_title(sub_items[section_no])
                    notes.doc.add_heading(h, level=2); notes.section_heading = h
                else:
                    review.append(f'{where}: agenda slide with no matching section name')
                section_no += 1
            elif agenda == 'first' and title_level == 3:
                level = 2                     # the opening agenda heads the slides before the first section
            hp = notes.doc.add_heading(slide_heading(n, n, tt), level=level)
            stats_inc('slide headings')
            if not tt: review.append(f'{where}: no slide title; the heading is the slide number alone')
            if agenda == 'section':
                run = None; continue          # its list repeats the opening agenda
        collect_figures(s)
        walk(s.shapes._spTree, where, t_el)
        flush_symbol()
        if notes.prefix and place_trailing_narration(where): pass
        elif notes.prefix:
            review.append(f'{where}: narration "{notes.prefix}" had nothing after it on the slide')
            p = notes.doc.add_paragraph(notes.prefix); notes.prefix = ''
        if SLIDE_HEADINGS:
            shared = notes.this_slide & notes.prev_slide
            if (BUILD_RANGES and agenda is None and run is not None and tt and run['last'] == n - 1
                    and norm_title(tt) == norm_title(run['title']) and shared):
                hp._p.getparent().remove(hp._p)           # a build: its additions go under one heading
                run['last'] = n; run['p'].text = slide_heading(run['first'], n, run['title'])
                stats_inc('slide headings', -1); stats_inc('build slides joined to the heading before')
            else:
                run = {'p': hp, 'first': n, 'last': n, 'title': tt} if agenda is None else None
    if R_OUTPUT_TABLES:
        notes.stats['R output tables'] = r_output_tables(notes.doc, notes.code_style_id)
    nb, ne = add_code_markers(notes.doc, notes.code_style_id)
    notes.stats['code blocks marked'] = nb; notes.stats['code end markers'] = ne
    notes.stats['empty paragraphs removed'] = tidy_spacing(notes.doc, notes.code_style_id)
    if NUMBER_RULE: number_pass(notes.doc, notes.code_style_id)
    if args.out:
        notes.doc.save(args.out); print('wrote', args.out)
    else:
        snap = os.path.join(os.path.dirname(os.path.abspath(args.deck)), '_snapshots')
        os.makedirs(snap, exist_ok=True)
        gen = os.path.join(snap, os.path.basename(base) + '_generated.docx')
        notes.doc.save(gen); print('wrote', gen)
        master = base + '.docx'
        if os.path.exists(master):
            print(f'NOT overwritten: {master} exists (the edited master). Compare with Review > Compare.')
        else:
            shutil.copyfile(gen, master); print('wrote', master)
    for k, v in notes.stats.items(): print(f'  {k}: {v}')
    seen, uniq = set(), []
    for r in review:
        k = r.split(': ', 1)[-1]
        if k not in seen: seen.add(k); uniq.append(r)
    print('review:' if uniq else 'review: nothing flagged')
    for r in uniq: print('  -', r)
    if notes.changes:
        print(f'numbers and symbols rewritten for listening, or left as possible years ({len(notes.changes)}; check each):')
        for w, kind, before, after in notes.changes: print(f'  - {w}, {kind}: {before} -> {after}')
    if args.changes:
        import csv
        with open(args.changes, 'w', newline='', encoding='utf-8') as fh:
            cw = csv.writer(fh); cw.writerow(['where', 'kind', 'before', 'after']); cw.writerows(notes.changes)
        print('wrote', args.changes)

if __name__ == '__main__':
    main()
